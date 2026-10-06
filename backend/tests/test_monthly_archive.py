from datetime import date

from openpyxl import load_workbook

from app.core.security import hash_password
from app.models.car import Car
from app.models.inventory import InventoryItem
from app.models.invoice import Invoice, InvoiceStatus
from app.models.service import Service
from app.models.tenant import Tenant
from app.models.user import Role, User
from app.services import monthly_archive_service


def test_monthly_archive_creates_excel_file(db, tmp_path, monkeypatch):
    monkeypatch.setattr(monthly_archive_service, "EXPORT_DIR", tmp_path)

    tenant = Tenant(name="Archive Center", plan="basic", contact_phone="07700000001")
    db.add(tenant)
    db.flush()
    manager = User(
        tenant_id=tenant.id,
        email="manager@archive.test",
        hashed_password=hash_password("pass"),
        role=Role.manager,
    )
    car = Car(tenant_id=tenant.id, plate_number="123ABC", owner_name="Ali", phone="07700000001")
    db.add_all([manager, car])
    db.flush()
    service = Service(
        tenant_id=tenant.id,
        car_id=car.id,
        service_date=date(2026, 5, 15),
        oil_type="تبديل زيت",
        mileage=10000,
    )
    db.add(service)
    db.flush()
    invoice = Invoice(
        tenant_id=tenant.id,
        service_id=service.id,
        amount=50000,
        discount=5000,
        status=InvoiceStatus.paid,
        invoice_date=date(2026, 5, 15),
    )
    item = InventoryItem(tenant_id=tenant.id, oil_type="15W40", quantity=4, min_threshold=2)
    db.add_all([invoice, item])
    db.commit()

    archive = monthly_archive_service.export_tenant_monthly_archive(db, tenant, 2026, 5)

    workbook = load_workbook(archive["path"])
    assert set(["Summary", "Cars", "Services", "Invoices", "Open Debts", "Inventory"]).issubset(workbook.sheetnames)
    assert workbook["Summary"]["B2"].value == "Archive Center"
    assert workbook["Services"]["E2"].value == "تبديل زيت"
    assert workbook["Invoices"]["G2"].value == 45000
    assert workbook["Inventory"]["B2"].value == "15W40"


def test_superadmin_can_run_monthly_archive(client, db, superadmin, superadmin_token, tmp_path, monkeypatch):
    monkeypatch.setattr(monthly_archive_service, "EXPORT_DIR", tmp_path)

    tenant = Tenant(name="Manual Archive", plan="enterprise", contact_phone="07700000002")
    db.add(tenant)
    db.commit()

    response = client.post(
        "/tenants/monthly-archives/run?year=2026&month=5",
        headers={"Authorization": f"Bearer {superadmin_token}"},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["tenant_count"] == 1
    assert data["results"][0]["tenant_name"] == "Manual Archive"
    assert data["results"][0]["whatsapp_status"] == "not_configured"
